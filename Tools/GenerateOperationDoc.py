#!/usr/bin/env python3
"""Generate the operation manual Word document for IamNotCardGod."""

from __future__ import annotations

import pathlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = pathlib.Path("我不是牌神 V1.0操作说明文档.docx")
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x59, 0x59, 0x59)

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

normal = doc.styles["Normal"]
normal.font.name = "SimSun"
normal.font.size = Pt(10.5)
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")


def set_font(run, size=10.5, bold=False, color=BLACK, ascii_font="SimSun", east="SimSun"):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_font(run, 10.5, False, BLACK)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(placeholder)
    run._r.append(end)


header = section.header
header.is_linked_to_previous = False
hp = header.paragraphs[0]
hp.text = ""
hp.paragraph_format.tab_stops.add_tab_stop(Cm(8), WD_TAB_ALIGNMENT.CENTER)
hp.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
hr = hp.add_run("\t我不是牌神V1.0\t")
set_font(hr, 10.5, False, BLACK)
add_page_field(hp)

settings = doc.settings.element
uf = OxmlElement("w:updateFields")
uf.set(qn("w:val"), "true")
settings.append(uf)


def para(text="", size=10.5, bold=False, color=BLACK, align=None, heading=False, before=0, after=2.3):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.5
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        if heading:
            set_font(r, size, True, BLACK, "Arial", "SimHei")
        else:
            set_font(r, size, bold, color)
    return p


def shot(text):
    para("【截图：" + text + "】", 10.5, True, GRAY, WD_ALIGN_PARAGRAPH.CENTER, False, 8, 10)


def add(kind, text=""):
    if kind == "title":
        para(text, 18, True, BLACK, WD_ALIGN_PARAGRAPH.CENTER, True, 0, 10)
    elif kind == "h1":
        para(text, 15, True, BLACK, None, True, 12, 5)
    elif kind == "h2":
        para(text, 13, True, BLACK, None, True, 10, 4)
    elif kind == "h3":
        para(text, 12, True, BLACK, None, True, 8, 4)
    elif kind == "para":
        para(text, 10.5, False, BLACK)
    elif kind == "bullet":
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(2.3)
        pf.line_spacing = 1.5
        pf.left_indent = Cm(0.6)
        r = p.add_run(text)
        set_font(r, 10.5, False, BLACK)
    elif kind == "num":
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(2.3)
        pf.line_spacing = 1.5
        pf.left_indent = Cm(0.6)
        r = p.add_run(text)
        set_font(r, 10.5, False, BLACK)
    elif kind == "shot":
        shot(text)
    elif kind == "note":
        para(text, 10, False, GRAY)


C = []
C.append(("title", "我不是牌神 V1.0操作说明文档"))
C.append(("note", "本文档根据软件实际功能编写，用于说明软件的操作流程、功能模块和界面使用方式。文中【截图：...】位置请在定稿时替换为对应界面截图。"))

C.append(("h1", "第一章 系统描述"))
C.append(("h2", "1.1 系统价值"))
C.append(("para", "本软件是一款面向抖音小游戏平台的休闲卡牌对战游戏，为玩家提供免下载、即点即玩的卡牌对战体验。玩家无需安装额外应用，即可在抖音客户端内随时进入游戏，体验策略卡牌对战的乐趣。"))
C.append(("para", "游戏提供智能AI对手、丰富的卡牌组合与PvE关卡挑战，满足玩家碎片化时间的娱乐需求。玩家可以在人机对战、玩家匹配和好友房等模式中选择适合的对战方式，并通过不断通关解锁更高级的卡牌。"))
C.append(("para", "软件在技术层面采用Unity引擎、C#语言和分层代码结构，结合HybridCLR代码热更新与YooAsset资源更新，能够在不重新发布小游戏包的情况下更新代码和资源，为后续玩法迭代提供支持。"))

C.append(("h2", "1.2 系统功能描述"))
C.append(("para", "本软件主要包含以下功能模块："))
C.append(("bullet", "游戏启动与场景管理：完成启动画面、资源初始化、登录、版本检查和场景跳转。"))
C.append(("bullet", "账号与数据管理：支持抖音平台登录、本地存档和云存档。"))
C.append(("bullet", "卡牌系统：包含基础手牌、强化卡牌、手牌上限、弃牌和合成规则。"))
C.append(("bullet", "战斗系统：包含回合流程、出牌、目标选择、战斗结算和胜负判定。"))
C.append(("bullet", "技能与效果系统：包含攻击、防御、治疗、增益、减益和持续状态效果。"))
C.append(("bullet", "AI与Boss系统：控制人机对手的决策、出牌延迟和Boss行为。"))
C.append(("bullet", "解锁与进度系统：记录PvE胜利、Boss通关和卡牌解锁进度。"))
C.append(("bullet", "热更新与资源更新：支持代码热更和资源包动态下载。"))
C.append(("bullet", "表现与反馈：提供动画、音效、飘字、特效和UI提示。"))

C.append(("h1", "第二章 系统模块介绍"))
C.append(("h2", "2.1 游戏启动"))
C.append(("para", "游戏启动后首先展示品牌启动画面，同时后台自动完成资源完整性校验、网络环境检测和版本号比对。若存在需要更新的资源或代码，系统会进入下载流程并显示更新进度。资源初始化完成后自动进入主界面。"))
C.append(("para", "启动阶段的完整操作流程如下："))
C.append(("num", "1. 在抖音客户端中进入本小游戏。"))
C.append(("num", "2. 等待品牌启动画面加载。"))
C.append(("num", "3. 系统自动初始化抖音登录、账号数据和资源服务。"))
C.append(("num", "4. 如检测到更新，自动下载最新资源并显示进度。"))
C.append(("num", "5. 初始化完成后进入游戏主界面。"))
C.append(("shot", "游戏启动画面"))
C.append(("shot", "资源更新界面"))

C.append(("h2", "2.2 游戏主界面"))
C.append(("para", "主界面为玩家提供游戏的核心入口。界面中展示玩家的头像、昵称、体力、段位分、卡牌解锁进度等基本信息，并提供开始战斗、玩家匹配、好友房、卡牌图鉴、每日福利、排行榜和设置等主要功能入口。"))
C.append(("para", "主界面的主要操作如下："))
C.append(("bullet", "点击“PVE”按钮进入普通人机对战。"))
C.append(("bullet", "点击“P2P”按钮进入玩家匹配对战。"))
C.append(("bullet", "点击“好友房”按钮创建好友房间或输入邀请码加入好友房间。"))
C.append(("bullet", "点击“图鉴/卡牌收藏”查看已解锁卡牌和Boss解锁进度。"))
C.append(("bullet", "点击“每日福利”领取体力、分享奖励和三场战斗奖励。"))
C.append(("bullet", "点击“排行榜”查看好友榜和世界榜。"))
C.append(("bullet", "点击“设置”调整音量等游戏设置。"))
C.append(("bullet", "点击“教程”重新进入新手引导。"))
C.append(("shot", "游戏主界面"))

C.append(("h2", "2.3 卡牌系统"))
C.append(("para", "卡牌系统是游戏的核心玩法基础。游戏中卡牌分为基础手牌和强化卡牌两类。基础手牌可以直接打出，强化卡牌需要拖拽合成到同类型的基础手牌上，以提升基础手牌的数值并附加额外效果。"))
C.append(("h3", "2.3.1 基础手牌"))
C.append(("bullet", "攻击牌：对敌方目标造成伤害，攻击牌会随回合进行成长。"))
C.append(("bullet", "防御牌：为玩家提供防御，防御牌可以在备战阶段提前打出。"))
C.append(("bullet", "治疗牌：为玩家恢复生命值，提高续航能力。"))
C.append(("para", "每名玩家在首回合会获得基础手牌和强化牌，后续回合系统会按照配置继续发牌。手牌超过上限时，玩家需要进入弃牌阶段弃掉多余卡牌。"))
C.append(("h3", "2.3.2 强化卡牌"))
C.append(("bullet", "强化卡牌分为银卡、金卡、白金卡和棱彩卡，品质越高，效果越强。"))
C.append(("bullet", "强化卡牌必须与基础手牌类型一致才能合成，例如攻击强化牌只能合成到攻击牌上。"))
C.append(("bullet", "每张基础手牌最多可以挂载两张强化卡牌。"))
C.append(("bullet", "部分防御强化牌可以在备战阶段使用，其他强化牌在战斗阶段使用。"))
C.append(("h3", "2.3.3 卡牌操作说明"))
C.append(("num", "1. 点击卡牌可以查看卡牌名称、类型、基础数值、成长、强化槽和效果说明。"))
C.append(("num", "2. 在战斗阶段双击基础手牌即可打出该卡牌。"))
C.append(("num", "3. 将强化卡牌拖拽到同类型基础手牌上，即可完成合成。"))
C.append(("num", "4. 在手牌超限时，点击弃牌按钮进入弃牌模式，再点击需要弃掉的卡牌。"))
C.append(("num", "5. 在卡牌图鉴中可以查看当前已解锁卡牌和后续解锁目标。"))
C.append(("shot", "卡牌图鉴/卡牌说明"))
C.append(("shot", "战斗中手牌管理"))
C.append(("shot", "强化卡合成"))

C.append(("h2", "2.4 战斗系统"))
C.append(("para", "战斗采用实时回合制规则。玩家和AI分别管理自己的手牌、牌库和弃牌堆，在备战阶段和战斗阶段完成出牌操作。战斗界面清晰展示双方生命值、防御、护盾、当前阶段和剩余时间。"))
C.append(("para", "一局战斗的完整流程如下："))
C.append(("num", "1. 玩家选择对局模式并消耗体力进入战斗场景。"))
C.append(("num", "2. 系统为双方发牌，首回合发放基础手牌和强化牌。"))
C.append(("num", "3. 如果手牌数量超过上限，进入弃牌阶段，玩家需要弃掉多余卡牌。"))
C.append(("num", "4. 进入备战阶段，玩家可以打出防御基础牌或使用备战期防御强化牌。"))
C.append(("num", "5. 点击准备按钮或等待备战倒计时结束后进入战斗阶段。"))
C.append(("num", "6. 在战斗阶段中，玩家可以合成强化牌、双击打出基础手牌，也可以选择结束本轮出牌。"))
C.append(("num", "7. 双方完成出牌或战斗倒计时结束后，进入回合结束，并开始下一回合。"))
C.append(("num", "8. 当任意一方生命值降至0时，战斗结束并进入结算界面。"))
C.append(("h3", "2.4.1 战斗属性"))
C.append(("bullet", "生命值：初始生命值为50，生命值降至0的一方判负。"))
C.append(("bullet", "防御：用于减少受到的伤害，不直接抵消伤害。"))
C.append(("bullet", "吸收护盾：优先吸收受到的伤害，护盾被击破时可触发特殊效果。"))
C.append(("bullet", "备战防御：在备战阶段打出防御基础牌后，本回合获得额外防御。"))
C.append(("bullet", "状态效果：包括灼烧、持续治疗、禁牌、圣盾、荆棘、反震、虚弱等。"))
C.append(("para", "玩家在出牌前需要根据当前生命、防御、护盾和手牌情况选择目标。攻击牌通常选择敌方目标，防御牌和治疗牌通常选择自己。系统会根据卡牌效果实时计算伤害、防御和治疗数值，并在界面中显示。"))
C.append(("shot", "战斗主界面"))
C.append(("shot", "出牌操作"))
C.append(("shot", "战斗结算界面"))

C.append(("h2", "2.5 技能与效果系统"))
C.append(("para", "技能与效果系统驱动每张卡牌释放后的具体表现。卡牌效果在战斗中通过特效动画和数值飘字直观呈现。"))
C.append(("bullet", "攻击效果：对目标造成伤害，并伴有打击特效和伤害飘字。"))
C.append(("bullet", "防御效果：为玩家添加防御或护盾，并在血条旁显示护盾图标。"))
C.append(("bullet", "治疗效果：恢复生命值，并伴有治疗特效。"))
C.append(("bullet", "增益效果：提升攻击力、防御力、治疗量或获得额外回合。"))
C.append(("bullet", "减益效果：降低敌方攻击、防御，或附加禁牌、灼烧等状态。"))
C.append(("bullet", "持续效果：灼烧和持续治疗等效果会在后续回合或时间内持续生效。"))
C.append(("para", "常见卡牌效果包括：蓄力、穿透、灼烧、吸血、猛击、精准、铁壁、护盾、反震、坚韧、圣盾、荆棘、吸收、反射、虚弱打击、处决、真实伤害、压制、碎甲、绝对防御、复仇、不屈、镜像、圣疗、救赎、生命窃取、逆转等。"))
C.append(("shot", "攻击效果"))
C.append(("shot", "防御护盾效果"))
C.append(("shot", "治疗效果"))

C.append(("h2", "2.6 AI与Boss系统"))
C.append(("para", "AI系统控制敌方对手的智能决策。在敌方回合中，AI根据当前手牌、生命值、防御和战场态势计算最优出牌策略，并模拟出牌延迟，提升对战的真实感。"))
C.append(("bullet", "AI会根据自身血量、护盾和手牌类型选择攻击、防御、治疗或合成操作。"))
C.append(("bullet", "AI具有出牌延迟和技能使用概率，避免行为过于机械。"))
C.append(("bullet", "Boss具有独立配置，包括额外生命值、出牌速度、攻击倾向、防御倾向和偏好卡牌。"))
C.append(("bullet", "Boss战界面显示Boss名称、生命值和当前挑战进度。"))
C.append(("para", "在PvE模式中，玩家每完成两场普通人机胜利，即可挑战下一位Boss。游戏共设置9个Boss，通关Boss后解锁对应的白金卡或棱彩卡强化牌。"))
C.append(("shot", "AI对战界面"))
C.append(("shot", "Boss战界面"))

C.append(("h2", "2.7 解锁与进度系统"))
C.append(("para", "进度系统驱动玩家的成长体验。界面中展示当前Boss进度、已解锁卡牌数量和任务完成状态。"))
C.append(("bullet", "普通PvE胜利会记录进度，满足条件后开放Boss挑战。"))
C.append(("bullet", "Boss通关后解锁对应的高级强化卡牌。"))
C.append(("bullet", "卡牌图鉴中会显示9个Boss对应的奖励卡牌是否已解锁。"))
C.append(("bullet", "主界面会显示下一目标提示，例如“普通人机胜利：1/2”或“下一目标：Boss 1/9”。"))
C.append(("bullet", "通关奖励可以在主界面的奖励面板中领取。"))
C.append(("shot", "关卡推进界面"))
C.append(("shot", "任务奖励"))

C.append(("h2", "2.8 热更新与资源更新"))
C.append(("para", "游戏启动或运行时，检测到新版本会自动触发更新流程。更新界面显示当前版本号、更新内容描述和下载进度条，确保玩家在资源下载过程中了解实时进度。"))
C.append(("bullet", "启动时自动检查远端资源版本清单。"))
C.append(("bullet", "存在新资源时自动下载并完成校验。"))
C.append(("bullet", "支持代码热更程序集加载，减少重新发版次数。"))
C.append(("bullet", "资源下载过程中显示进度，下载完成后继续进入游戏。"))
C.append(("shot", "资源更新界面"))

C.append(("h2", "2.9 表现与反馈"))
C.append(("para", "游戏中所有操作均配有即时反馈。卡牌释放时播放对应动画和音效，伤害和治疗数值以飘字形式在目标头顶显示，状态变化通过图标闪烁或文字提示告知玩家。"))
C.append(("bullet", "发牌、出牌、合成、弃牌等操作具有动画反馈。"))
C.append(("bullet", "攻击、防御、治疗和护盾效果具有对应特效。"))
C.append(("bullet", "战斗阶段和备战阶段显示倒计时。"))
C.append(("bullet", "按钮点击、弹窗提示和场景跳转具有UI反馈。"))
C.append(("bullet", "音效系统提供背景音乐、卡牌音效和战斗音效。"))
C.append(("shot", "战斗动画"))
C.append(("shot", "UI提示"))
C.append(("shot", "特效反馈"))

C.append(("h1", "第三章 操作说明汇总"))
C.append(("para", "为了方便玩家快速上手，下面汇总游戏中的常见操作方式。"))
C.append(("bullet", "进入对战：在主界面点击PVE、P2P或好友房入口。"))
C.append(("bullet", "查看卡牌：点击手牌卡牌或打开图鉴。"))
C.append(("bullet", "合成强化牌：将强化牌拖拽到同类型基础牌上。"))
C.append(("bullet", "打出基础牌：在战斗阶段双击基础牌。"))
C.append(("bullet", "弃牌：点击弃牌按钮进入弃牌模式，再点击多余手牌。"))
C.append(("bullet", "结束备战：点击准备按钮。"))
C.append(("bullet", "结束出牌：在战斗阶段点击跳过本轮战斗。"))
C.append(("bullet", "查看进度：在主界面查看体力、段位分和Boss解锁进度。"))
C.append(("bullet", "领取福利：打开每日福利面板领取奖励。"))
C.append(("bullet", "调整设置：打开设置面板调整音量等选项。"))

C.append(("h1", "第四章 体力与福利"))
C.append(("para", "体力是进入对战所消耗的资源，游戏通过体力机制控制玩家对局频率，并提供多种恢复方式。"))
C.append(("bullet", "体力上限为8点。"))
C.append(("bullet", "普通PvE对战消耗1点体力，Boss对战消耗2点体力。"))
C.append(("bullet", "每日自动恢复一定体力。"))
C.append(("bullet", "观看激励视频可以补充体力。"))
C.append(("bullet", "分享游戏可以领取分享奖励。"))
C.append(("bullet", "每日完成三场战斗后可以领取额外体力奖励。"))
C.append(("shot", "每日福利/体力界面"))

C.append(("h1", "第五章 账号与数据"))
C.append(("para", "游戏支持抖音平台账号登录，并保存玩家数据。"))
C.append(("bullet", "玩家可以通过抖音登录获取头像、昵称和用户信息。"))
C.append(("bullet", "游戏数据支持本地存档和云存档。"))
C.append(("bullet", "保存内容包括卡牌解锁进度、Boss进度、体力、玩家昵称和段位分等。"))
C.append(("bullet", "玩家可以在个人信息界面修改昵称和选择头像。"))
C.append(("bullet", "排行榜数据来自玩家对战段位分。"))
C.append(("shot", "登录/个人信息界面"))

C.append(("h1", "第六章 设置与帮助"))
C.append(("para", "设置界面提供基础游戏参数调整，教程和图鉴功能帮助玩家快速了解玩法。"))
C.append(("bullet", "音量设置：调整背景音乐和音效音量。"))
C.append(("bullet", "新手教程：重新进入分步操作引导。"))
C.append(("bullet", "卡牌图鉴：查看全部卡牌和解锁状态。"))
C.append(("bullet", "返回主界面：战斗中可以返回主界面。"))
C.append(("bullet", "兑换码：在兑换码面板输入兑换码领取奖励。"))
C.append(("shot", "设置界面"))

for kind, text in C:
    add(kind, text)

doc.core_properties.title = "我不是牌神 V1.0操作说明文档"
doc.save(OUT)
print("saved", OUT.resolve())
print("paragraphs", len(doc.paragraphs))
