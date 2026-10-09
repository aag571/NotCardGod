#!/usr/bin/env python3
"""Generate the software copyright source-code document for IamNotCardGod."""

from __future__ import annotations

import pathlib
import unicodedata

from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = pathlib.Path(__file__).resolve().parents[1]
SRC = ROOT / "Assets" / "Scripts"
OUT = ROOT / "\u6211\u4e0d\u662f\u724c\u795ev1.0-\u6e90\u4ee3\u7801\u6587\u6863.docx"

TITLE = "\u6211\u4e0d\u662f\u724c\u795ev1.0"
LINES_PER_PART = 1500
TOTAL_LINES = LINES_PER_PART * 2
MAX_LINE_WIDTH_PT = 460


FIRST_FILES = [
    "AOT/Animation/AnimationQueue.cs",
    "AOT/Core/GameApp.cs",
    "AOT/Core/GameServiceRegistry.cs",
    "AOT/Core/SceneContext.cs",
    "AOT/Data/Configs/CardArtConfig.cs",
    "AOT/Data/GameDataRepository.cs",
    "AOT/Gameplay/Cards/BasicHandCard.cs",
    "AOT/Gameplay/Cards/StrengthenedSkillCard.cs",
    "AOT/Gameplay/Combat/CardMaven.cs",
    "AOT/Gameplay/Combat/Desk.cs",
    "AOT/Progression/LocalAccountService.cs",
    "AOT/Core/GameController.cs",
]

SECOND_FILES = [
    "AOT/Data/Configs/BossConfig.cs",
    "AOT/HotUpdate/SkillHot.cs",
    "AOT/Scene/LoadingSceneController.cs",
    "AOT/UI/BattleCardView.cs",
    "HotUpdate/CardManager.cs",
]


def strip_comments(text: str) -> str:
    """Remove C# line and block comments while preserving source line breaks."""
    result: list[str] = []
    i = 0
    n = len(text)

    while i < n:
        char = text[i]

        if char == "/" and i + 1 < n and text[i + 1] == "/":
            newline = text.find("\n", i)
            if newline == -1:
                i = n
            else:
                result.append("\n")
                i = newline + 1
            continue

        if char == "/" and i + 1 < n and text[i + 1] == "*":
            depth = 1
            i += 2
            while i < n and depth:
                if i + 1 < n and text[i] == "/" and text[i + 1] == "*":
                    depth += 1
                    i += 2
                elif i + 1 < n and text[i] == "*" and text[i + 1] == "/":
                    depth -= 1
                    i += 2
                else:
                    if text[i] == "\n":
                        result.append("\n")
                    i += 1
            continue

        if char == '"':
            verbatim = i > 0 and text[i - 1] == "@"
            result.append(char)
            i += 1
            while i < n:
                if verbatim:
                    if text[i] == '"':
                        if i + 1 < n and text[i + 1] == '"':
                            result.append('""')
                            i += 2
                        else:
                            result.append('"')
                            i += 1
                            break
                    else:
                        result.append(text[i])
                        i += 1
                else:
                    if text[i] == "\\":
                        result.append(text[i])
                        i += 1
                        if i < n:
                            result.append(text[i])
                            i += 1
                    elif text[i] == '"':
                        result.append(text[i])
                        i += 1
                        break
                    else:
                        result.append(text[i])
                        i += 1
            continue

        if char == "'":
            result.append(char)
            i += 1
            while i < n:
                if text[i] == "\\":
                    result.append(text[i])
                    i += 1
                    if i < n:
                        result.append(text[i])
                        i += 1
                elif text[i] == "'":
                    result.append(text[i])
                    i += 1
                    break
                else:
                    result.append(text[i])
                    i += 1
            continue

        result.append(char)
        i += 1

    return "".join(result)


def estimate_width(line: str) -> int:
    """Approximate display width in Word points for SimSun 10pt."""
    width = 0
    for char in line:
        if char == "\t":
            width += 20
        elif unicodedata.east_asian_width(char) in ("F", "W"):
            width += 10
        else:
            width += 5
    return width


def safe_break_positions(line: str) -> list[int]:
    """Return whitespace positions outside string/char literals."""
    breaks: list[int] = []
    in_string = False
    quote = ""
    escaped = False

    for i, char in enumerate(line):
        if in_string:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == quote:
                in_string = False
        else:
            if char in ('"', "'"):
                in_string = True
                quote = char
            elif char.isspace():
                breaks.append(i)

    return breaks


def wrap_line(line: str, max_width: int) -> list[str]:
    """Wrap one source line at whitespace outside literals when necessary."""
    line = line.replace("\t", "    ")
    if estimate_width(line) <= max_width:
        return [line]

    indent = line[: len(line) - len(line.lstrip())]
    continuation = indent + "    "
    breaks = safe_break_positions(line)
    pieces: list[str] = []
    position = 0

    while position < len(line):
        chosen = None
        for break_at in breaks:
            if break_at > position and estimate_width(line[position:break_at]) <= max_width:
                chosen = break_at

        if chosen is None:
            next_break = next((b for b in breaks if b > position), len(line))
            chosen = next_break

        piece = line[position:chosen].rstrip()
        if piece:
            pieces.append(piece)
        position = chosen
        while position < len(line) and line[position].isspace():
            position += 1

    result = [pieces[0]]
    for piece in pieces[1:]:
        result.append(continuation + piece.lstrip())
    return result


def build_part(files: list[str]) -> list[str]:
    """Build one 1500-line half with no blank lines."""
    display_lines: list[str] = []

    for relative in files:
        source_path = SRC / relative
        source = source_path.read_text(encoding="utf-8-sig", errors="replace")
        cleaned = strip_comments(source).splitlines()
        for line in cleaned:
            if line.strip():
                display_lines.extend(wrap_line(line, MAX_LINE_WIDTH_PT))

    if len(display_lines) != LINES_PER_PART:
        raise RuntimeError(
            f"Expected {LINES_PER_PART} lines, got {len(display_lines)}."
        )

    return display_lines


def set_run_font(run) -> None:
    """Set SimSun 10pt for both Latin and East Asian characters."""
    run.font.name = "SimSun"
    run.font.size = Pt(10)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "SimSun")
    rfonts.set(qn("w:hAnsi"), "SimSun")
    rfonts.set(qn("w:eastAsia"), "SimSun")
    rfonts.set(qn("w:cs"), "SimSun")


def add_page_number_field(paragraph) -> None:
    """Insert a Word PAGE field at the current run position."""
    run = paragraph.add_run()
    set_run_font(run)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(placeholder)
    run._r.append(end)


def set_paragraph_format(paragraph) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(15)

    ppr = paragraph._p.get_or_add_pPr()
    snap_to_grid = OxmlElement("w:snapToGrid")
    snap_to_grid.set(qn("w:val"), "0")
    ppr.append(snap_to_grid)


def configure_header(section) -> None:
    header = section.header
    header.is_linked_to_previous = False

    paragraph = header.paragraphs[0]
    paragraph.text = ""
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.tab_stops.add_tab_stop(Cm(8.5), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT)

    run = paragraph.add_run("\t" + TITLE + "\t")
    set_run_font(run)
    add_page_number_field(paragraph)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(1)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(1)

    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(10)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "SimSun")
    rfonts.set(qn("w:hAnsi"), "SimSun")
    rfonts.set(qn("w:eastAsia"), "SimSun")
    rfonts.set(qn("w:cs"), "SimSun")

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    configure_header(section)


def main() -> None:
    doc = Document()
    configure_document(doc)

    first_part = build_part(FIRST_FILES)
    second_part = build_part(SECOND_FILES)
    all_lines = first_part + second_part

    for index, line in enumerate(all_lines):
        paragraph = doc.add_paragraph()
        set_paragraph_format(paragraph)
        if index == LINES_PER_PART:
            paragraph.paragraph_format.page_break_before = True
        paragraph.add_run(line)
        set_run_font(paragraph.runs[0])

    doc.save(OUT)

    print(f"Output: {OUT}")
    print(f"Total source lines: {len(all_lines)}")
    print(f"First half lines: {len(first_part)}")
    print(f"Second half lines: {len(second_part)}")
    print(f"Blank lines: {sum(1 for line in all_lines if not line.strip())}")
    print(f"First half last line: {first_part[-1]!r}")
    print(f"Second half last line: {second_part[-1]!r}")


if __name__ == "__main__":
    main()